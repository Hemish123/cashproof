import pandas as pd
import os

def create_sample_files():
    # 1. Generate chase_checking.csv
    checking_data = [
        ["Chase Bank - Statement of Account"],
        ["Account Number: 1234"],
        ["Account Holder: John Doe"],
        ["Statement Period: 2026-08-01 to 2026-08-31"],
        [""],
        ["Date", "Description", "Amount", "Balance"],
        ["2026-08-01", "Opening Balance Forward", "1000.00", "1000.00"],
        ["2026-08-05", "Employer Payroll Direct Deposit", "2500.00", "3500.00"],
        ["2026-08-10", "Whole Foods Market Grocery", "-120.50", "3379.50"],
        ["2026-08-15", "Online Transfer to acct 5678", "-500.00", "2879.50"],
        ["2026-08-20", "Apartment Rent Payment", "-1000.00", "1879.50"],
        ["2026-08-28", "Starbucks Coffee", "-6.50", "1873.00"]
    ]
    
    checking_df = pd.DataFrame(checking_data)
    checking_df.to_csv("chase_checking.csv", index=False, header=False)
    print("Created chase_checking.csv")

    # 2. Generate chase_savings.xlsx
    savings_metadata = [
        ["Chase Savings Account Summary"],
        ["Account Number: 5678"],
        ["Account Holder: John Doe"],
        ["Statement Period: 2026-08-01 to 2026-08-31"],
        [""]
    ]
    
    savings_headers = ["Transaction Date", "Details", "Deposit", "Withdrawal", "Running Balance"]
    
    savings_rows = [
        ["2026-08-01", "Opening Balance", "5000.00", "", "5000.00"],
        ["2026-08-16", "Online Transfer from checking", "500.00", "", "5500.00"],
        ["2026-08-18", "Monthly Interest Earned", "5.50", "", "5505.50"],
        ["2026-08-25", "Zelle to Checking self", "", "100.00", "5405.50"]
    ]

    # Combine into Excel using openpyxl
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Savings Statement"
    
    # Write metadata
    for r_idx, row in enumerate(savings_metadata, start=1):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
            
    # Write headers at row 6
    for c_idx, val in enumerate(savings_headers, start=1):
        ws.cell(row=6, column=c_idx, value=val)
        
    # Write rows
    for r_idx, row in enumerate(savings_rows, start=7):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
            
    wb.save("chase_savings.xlsx")
    print("Created chase_savings.xlsx")

if __name__ == "__main__":
    create_sample_files()
