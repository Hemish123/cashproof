"""
cashproof_ai/reports.py

Generates an Excel workbook that exactly matches the QNB_KwikGoal_Cashflow_Report.xlsx format:

  Sheet 1: Executive Summary
  Sheet 2: Transaction Ledger
  Sheet 3…N: One per uploaded statement (named {Bank}_{AcctNum})

Styles mirror the reference file exactly:
  - Font: Segoe UI
  - Header fill colours: 1F4E79 (summary) / 2F5597 (ledger)
  - Currency: #,##0.00;(#,##0.00);"-";@
  - Zebra: F2F2F2
  - Total row: DDEBF7 with double-bottom border
  - Interbank rows: FFF2CC (amber/yellow)
"""

import re
import logging
from decimal import Decimal
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import CashProofSession, LedgerTransaction, ParsedStatement

logger = logging.getLogger(__name__)


# ── Style constants (matching QNB reference exactly) ─────────────────────────

FONT_SEGOE = 'Segoe UI'

font_title    = Font(name=FONT_SEGOE, size=16, bold=True,  color='1F4E79')
font_subtitle = Font(name=FONT_SEGOE, size=10, italic=True, color='595959')
font_header   = Font(name=FONT_SEGOE, size=11, bold=True,  color='FFFFFF')
font_body     = Font(name=FONT_SEGOE, size=11)
font_total    = Font(name=FONT_SEGOE, size=11, bold=True)
font_pass     = Font(name=FONT_SEGOE, size=11, bold=True,  color='2E7D32')
font_fail     = Font(name=FONT_SEGOE, size=11, bold=True,  color='C62828')

fill_header_summary = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
fill_header_ledger  = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
fill_zebra          = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
fill_total          = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')
fill_interbank      = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')

_thin      = Side(style='thin',   color='D9D9D9')
_blk_thin  = Side(style='thin',   color='000000')
_blk_dbl   = Side(style='double', color='000000')
border_all   = Border(left=_thin,     right=_thin,     top=_thin,     bottom=_thin)
border_total = Border(left=_thin,     right=_thin,     top=_blk_thin, bottom=_blk_dbl)

align_center = Alignment(horizontal='center', vertical='center')
align_left   = Alignment(horizontal='left',   vertical='center')
align_right  = Alignment(horizontal='right',  vertical='center')

NUM_CURRENCY = '#,##0.00;(#,##0.00);"-";@'


# ── Helpers ───────────────────────────────────────────────────────────────────

def _fmt_period(start, end):
    if start and end:
        return f'{start.strftime("%b %d")} - {end.strftime("%b %d")}, {start.year}'
    return 'Unknown'


def _fmt_statement_dates(start, end):
    if start and end:
        return f'Statement Dates {start.strftime("%-m/%d/%y")} thru {end.strftime("%-m/%d/%y")}'
    return 'Statement Dates Unknown'


def _fmt_statement_dates_safe(start, end):
    """Cross-platform version of _fmt_statement_dates (no %-m on Windows)."""
    if start and end:
        return f'Statement Dates {start.month}/{start.strftime("%d/%y")} thru {end.month}/{end.strftime("%d/%y")}'
    return 'Statement Dates Unknown'


def _safe_sheet_name(bank: str, acct: str) -> str:
    """Build a sheet name like QNB_9625, max 31 chars, no illegal chars."""
    bank_part = re.sub(r'[^\w]', '', (bank or 'Bank').split()[0])[:10]
    acct_part = re.sub(r'[^\w]', '', (acct or 'Acct'))[-8:]
    raw = f'{bank_part}_{acct_part}'
    raw = re.sub(r'[\\/*?:\[\]]', '', raw)
    return raw[:31]


def _autofit(ws, skip_rows=3, min_width=14, max_width=55):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col[skip_rows:]:
            if cell.value is not None:
                val_str = str(cell.value)
                if isinstance(cell.value, (float, int, Decimal)):
                    val_str = f'{cell.value:,.2f}'
                max_len = max(max_len, len(val_str))
        # Include header in width
        hdr = col[skip_rows].value if len(col) > skip_rows else None
        max_len = max(max_len, len(str(hdr or '')) + 3)
        min_w = 28 if col_letter == 'A' else min_width
        ws.column_dimensions[col_letter].width = min(max(max_len, min_w), max_width)


# ── Interbank status helpers ──────────────────────────────────────────────────

CONFIRMED = 'CONFIRMED_INTERBANK'

def _is_interbank(tx: LedgerTransaction) -> bool:
    return tx.interbank_status == CONFIRMED

def _tx_category(tx: LedgerTransaction) -> str:
    """Map §4 type code to a human-readable category string matching the QNB reference style."""
    labels = {
        'A': 'Customer Receipts / Revenue Credits',
        'B': 'Operating Payments / Vendor Disbursements',
        'C': 'Interbank Self-Transfer',
        'D': 'Interest Income',
        'E': 'Bank Fee / Service Charge',
        'F': 'Check',
        'G': 'ACH',
        'H': 'Card / Merchant Settlement',
        'I': 'Other / Unclassified',
    }
    return labels.get(tx.transaction_type, 'Other / Unclassified')


def _tx_net_amount(tx: LedgerTransaction) -> Decimal:
    """Return signed amount: credit positive, debit negative — matches QNB Amount column."""
    if tx.credit_amount and tx.credit_amount > 0:
        return tx.credit_amount
    if tx.debit_amount and tx.debit_amount > 0:
        return -tx.debit_amount
    return tx.net_amount or Decimal('0')


# ── Main entry point ──────────────────────────────────────────────────────────

def generate_report(session: CashProofSession, output_path: str) -> None:
    """
    Generate the Excel workbook matching QNB_KwikGoal_Cashflow_Report.xlsx format.

    Sheet order:
      1. Executive Summary
      2. Transaction Ledger
      3…N. One individual account sheet per ParsedStatement
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default blank sheet

    accounts  = list(session.client_accounts.all().order_by('bank_name', 'account_number'))
    statements = list(session.statements.filter(parse_status='COMPLETED').order_by('uploaded_at'))

    # Collect all transactions sorted newest-first (matching QNB ledger order)
    stmt_ids = [s.id for s in statements]
    all_txs = list(
        LedgerTransaction.objects.filter(statement_id__in=stmt_ids)
                                  .select_related('client_account', 'interbank_match')
                                  .order_by('-transaction_date', '-id')
    )

    company_name = session.client_name or 'Client'
    primary_bank = accounts[0].bank_name if accounts else 'Unknown Bank'

    # ── Build sheets ──────────────────────────────────────────────────────────
    ws_summary = wb.create_sheet('Executive Summary')
    ws_ledger  = wb.create_sheet('Transaction Ledger')

    _build_executive_summary(ws_summary, session, accounts, statements, company_name, primary_bank)
    _build_transaction_ledger(ws_ledger, all_txs, accounts, company_name, primary_bank)

    # One sheet per uploaded statement (matching QNB_9625 / QNB_4940 pattern)
    for stmt in statements:
        sheet_name = _safe_sheet_name(
            stmt.bank_name_extracted or primary_bank,
            stmt.account_number_extracted or 'Acct',
        )
        # Avoid duplicate sheet names if two stmts map to the same name
        base = sheet_name
        n = 2
        while base in [ws.title for ws in wb.worksheets]:
            base = f'{sheet_name}_{n}'
            n += 1
        ws_acc = wb.create_sheet(base)
        _build_account_sheet(ws_acc, stmt, company_name)

    wb.save(output_path)
    logger.info(f'Saved Excel report to {output_path}')


# ── Sheet 1: Executive Summary ────────────────────────────────────────────────

def _build_executive_summary(ws, session, accounts, statements, company_name, primary_bank):
    """Exact match to QNB_KwikGoal Executive Summary layout."""

    # Row 1 — title
    ws['A1'] = 'Bank Statement Cash Flow Analysis Report'
    ws['A1'].font = font_title

    # Row 2 — subtitle
    ws['A2'] = f'Executive Aggregates & Standardization Summary — {company_name} ({primary_bank})'
    ws['A2'].font = font_subtitle

    # Row 4 — column headers (same 9 as QNB)
    headers = [
        'Account', 'Bank Name', 'Account Number',
        'Total Deposits', 'Total Payments',
        'Total Interbank Deposits', 'Total Interbank Payments',
        'Net Deposits', 'Net Payments',
    ]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font      = font_header
        cell.fill      = fill_header_summary
        cell.alignment = align_center
        cell.border    = border_all
    ws.row_dimensions[4].height = 28

    # Data rows — one per ClientAccount
    current_row = 5
    totals = dict(deposits=Decimal('0'), payments=Decimal('0'),
                  ib_deposits=Decimal('0'), ib_payments=Decimal('0'),
                  net_deposits=Decimal('0'), net_payments=Decimal('0'))

    for idx, acc in enumerate(accounts):
        row_data = [
            acc.account_name or 'Operating Account',
            acc.bank_name,
            acc.account_number,
            float(acc.total_deposits),
            float(acc.total_payments),
            float(acc.total_interbank_deposits),
            float(acc.total_interbank_payments),
            float(acc.net_deposits),
            float(acc.net_payments),
        ]
        totals['deposits']    += acc.total_deposits
        totals['payments']    += acc.total_payments
        totals['ib_deposits'] += acc.total_interbank_deposits
        totals['ib_payments'] += acc.total_interbank_payments
        totals['net_deposits'] += acc.net_deposits
        totals['net_payments'] += acc.net_payments

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font   = font_body
            cell.border = border_all
            if col_idx <= 3:
                cell.alignment = align_left
            else:
                cell.alignment    = align_right
                cell.number_format = NUM_CURRENCY
            if idx % 2 == 1:
                cell.fill = fill_zebra

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # Combined Total row
    ws.cell(current_row, 1, 'Combined Total').font      = font_total
    ws.cell(current_row, 1).alignment                   = align_left
    ws.cell(current_row, 1).border                      = border_total
    ws.cell(current_row, 1).fill                        = fill_total
    ws.cell(current_row, 2, '').border                  = border_total
    ws.cell(current_row, 2).fill                        = fill_total
    ws.cell(current_row, 3, '').border                  = border_total
    ws.cell(current_row, 3).fill                        = fill_total

    for i, key in enumerate(['deposits', 'payments', 'ib_deposits', 'ib_payments', 'net_deposits', 'net_payments'], start=4):
        cell = ws.cell(current_row, i, float(totals[key]))
        cell.font          = font_total
        cell.alignment     = align_right
        cell.number_format = NUM_CURRENCY
        cell.border        = border_total
        cell.fill          = fill_total
    ws.row_dimensions[current_row].height = 22

    # Source statement reconciliation block (one per statement)
    current_row += 2
    for stmt in statements:
        beg = stmt.beginning_balance or Decimal('0')
        end = stmt.ending_balance or Decimal('0')
        stmt_txs = stmt.transactions.all()

        total_cred = sum((t.credit_amount or Decimal('0')) for t in stmt_txs if t.credit_amount)
        total_deb  = sum((t.debit_amount  or Decimal('0')) for t in stmt_txs if t.debit_amount)
        calc_end   = beg + total_cred - total_deb
        diff       = abs(calc_end - end)
        recon_ok   = diff <= Decimal('0.05')
        recon_txt  = ('PASS' if recon_ok else 'FAIL') + ' — Beginning + Credits - Debits = Ending'

        period_str = _fmt_period(stmt.period_start, stmt.period_end)
        bank_str   = stmt.bank_name_extracted or 'Unknown Bank'

        def _lbl(r, text):
            ws.cell(r, 1, text).font = font_total

        def _val(r, val, is_num=True):
            cell = ws.cell(r, 2, val)
            cell.font = font_body
            if is_num:
                cell.number_format = NUM_CURRENCY
                cell.alignment     = align_right

        _lbl(current_row, 'Source Statement')
        ws.cell(current_row, 2, f'{bank_str} — {period_str}').font = font_body
        current_row += 1

        _lbl(current_row, 'Beginning Balance');   _val(current_row, float(beg));        current_row += 1
        _lbl(current_row, 'Ending Balance');      _val(current_row, float(end));        current_row += 1
        _lbl(current_row, 'Statement Debits');    _val(current_row, float(total_deb));  current_row += 1
        _lbl(current_row, 'Statement Credits');   _val(current_row, float(total_cred)); current_row += 1

        ws.cell(current_row, 1, 'Reconciliation Check').font = font_total
        ws.cell(current_row, 2, recon_txt).font = font_pass if recon_ok else font_fail
        current_row += 2  # gap

    _autofit(ws)


# ── Sheet 2: Transaction Ledger ───────────────────────────────────────────────

def _build_transaction_ledger(ws, all_txs, accounts, company_name, primary_bank):
    """Exact match to QNB_KwikGoal Transaction Ledger layout."""

    # Build account summary string like "9625 Operating, 4940 Payroll, 7323 Escrow MM"
    account_summaries = []
    for acc in accounts:
        title = (acc.account_name or 'Operating').replace(' Account', '').strip()
        account_summaries.append(f'{acc.account_number} {title}')
    accounts_str = ', '.join(account_summaries) if account_summaries else 'All Accounts'

    ws['A1'] = 'Standardized Transactions Ledger'
    ws['A1'].font = font_title

    ws['A2'] = (f'Combined granular transaction logs across all {primary_bank} accounts '
                f'({accounts_str}) — {company_name}')
    ws['A2'].font = font_subtitle

    # Row 4 — headers (exact 9 columns matching QNB reference)
    ledger_headers = [
        'Date', 'Bank Name', 'Account Number', 'Account Title',
        'Statement Ref', 'Description', 'Amount', 'Category', 'Interbank Self-Transfer?',
    ]
    for col_idx, h in enumerate(ledger_headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font      = font_header
        cell.fill      = fill_header_ledger
        cell.alignment = align_center
        cell.border    = border_all
    ws.row_dimensions[4].height = 28

    current_row = 5
    for idx, tx in enumerate(all_txs):
        # Statement Ref: e.g. "9625 - Jan 2026"
        month_abbr = tx.transaction_date.strftime('%b')
        stmt_ref   = f'{tx.account_number} - {month_abbr} {tx.transaction_date.year}'

        amount    = _tx_net_amount(tx)
        is_ib     = _is_interbank(tx)
        category  = _tx_category(tx)

        row_data = [
            tx.transaction_date.strftime('%Y-%m-%d'),
            tx.bank_name,
            tx.account_number,
            tx.account_title or 'Operating Account',
            stmt_ref,
            tx.description,
            float(amount),
            category,
            'Yes' if is_ib else 'No',
        ]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font   = font_body
            cell.border = border_all

            if col_idx in (1, 9):
                cell.alignment = align_center
            elif col_idx == 7:
                cell.alignment     = align_right
                cell.number_format = NUM_CURRENCY
            else:
                cell.alignment = align_left

            # Interbank rows: amber fill; "Yes" cell bold
            if is_ib:
                cell.fill = fill_interbank
                if col_idx == 9:
                    cell.font = font_total
            else:
                if idx % 2 == 1:
                    cell.fill = fill_zebra

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    _autofit(ws)


# ── Sheet 3…N: Individual account sheets (one per ParsedStatement) ───────────

def _build_account_sheet(ws, stmt: ParsedStatement, company_name: str):
    """
    Exact match to QNB_9625 / QNB_4940 / QNB_7323 sheet layout.
    Columns: A = label, B = value — no data grid, just a metadata card + per-account detail.
    """
    bank_name  = stmt.bank_name_extracted or 'Unknown Bank'
    acct_num   = stmt.account_number_extracted or 'Unknown'
    acct_title = stmt.account_title_extracted or 'Operating Account'
    period_str = _fmt_statement_dates_safe(stmt.period_start, stmt.period_end)

    txs = list(stmt.transactions.all().order_by('transaction_date', 'id'))

    beg = stmt.beginning_balance or Decimal('0')
    end = stmt.ending_balance    or Decimal('0')

    total_cred = sum((t.credit_amount or Decimal('0')) for t in txs if t.credit_amount)
    total_deb  = sum((t.debit_amount  or Decimal('0')) for t in txs if t.debit_amount)
    calc_end   = beg + total_cred - total_deb
    diff       = abs(calc_end - end)
    recon_ok   = diff <= Decimal('0.05')
    recon_status = 'PASS' if recon_ok else 'FAIL'

    # Row 1 header — "Bank Statement" | bank_name  (matching QNB reference)
    cell_h1 = ws.cell(1, 1, 'Bank Statement')
    cell_h1.font      = font_header
    cell_h1.fill      = fill_header_summary
    cell_h1.alignment = align_left
    cell_h1.border    = border_all

    cell_h2 = ws.cell(1, 2, bank_name)
    cell_h2.font      = font_header
    cell_h2.fill      = fill_header_summary
    cell_h2.alignment = align_left
    cell_h2.border    = border_all

    # Metadata card (matching exact QNB row order)
    card = [
        ('Account',                   acct_num),
        ('Account Title',             acct_title),
        ('Statement Period',          period_str),
        ('Transactions in Ledger',    len(txs)),
        ('Statement Total Credits',   float(total_cred)),
        ('FDD Ledger Total Deposits', float(total_cred)),
        ('Statement Total Debits',    float(total_deb)),
        ('FDD Ledger Total Payments', float(total_deb)),
        ('Statement Beginning Balance', float(beg)),
        ('Statement Ending Balance',  float(end)),
        ('Calculated Ending Balance', float(calc_end)),
        ('Reconciliation',            recon_status),
    ]

    for row_offset, (label, val) in enumerate(card, start=2):
        cell_lbl = ws.cell(row_offset, 1, label)
        cell_lbl.font      = font_total
        cell_lbl.alignment = align_left
        cell_lbl.border    = border_all

        cell_val = ws.cell(row_offset, 2, val)
        cell_val.font   = font_body
        cell_val.border = border_all

        if isinstance(val, float) and label != 'Transactions in Ledger':
            cell_val.alignment     = align_right
            cell_val.number_format = NUM_CURRENCY
        elif label == 'Transactions in Ledger':
            cell_val.alignment     = align_right
            cell_val.number_format = '#,##0'
        else:
            cell_val.alignment = align_left

        if label == 'Reconciliation':
            cell_val.font = font_pass if recon_ok else font_fail

    # Autofit the two columns
    max_a = max(len(str(row[0])) for row in card) + 3
    max_b = 20
    ws.column_dimensions['A'].width = max(max_a, 32)
    ws.column_dimensions['B'].width = max(max_b, 22)
